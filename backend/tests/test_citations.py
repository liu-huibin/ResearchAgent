import hashlib
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from docx import Document as WordDocument
from PyPDF2 import PdfWriter
from PyPDF2.generic import DictionaryObject, NameObject, DecodedStreamObject

from app.models.document import Document
from app.services.citations import locate_chunk, read_units, resolve_citation
from app.services.chunking import chunk_document


def make_pdf(path):
    writer = PdfWriter()
    for text in ('first page alpha', 'second page beta'):
        writer.add_blank_page(width=600, height=800)
        page = writer.pages[-1]
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                                 NameObject('/Subtype'): NameObject('/Type1'),
                                 NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
        stream = DecodedStreamObject()
        stream.set_data(f'BT /F1 12 Tf 40 700 Td ({text}) Tj ET'.encode())
        page[NameObject('/Contents')] = writer._add_object(stream)
    with path.open('wb') as output:
        writer.write(output)


def make_browser_fixtures(directory):
    """Synthetic documents only; shared by the browser regression runner."""
    root = Path(directory)
    root.mkdir(parents=True, exist_ok=True)
    make_pdf(root / 'source.pdf')
    doc = WordDocument()
    for i in range(35):
        doc.add_paragraph(f'Introduction paragraph {i}')
    paragraph = doc.add_paragraph()
    paragraph.add_run('target ').bold = True
    paragraph.add_run('paragraph')
    doc.add_paragraph('second target')
    doc.save(root / 'source.docx')


class CitationTests(unittest.TestCase):
    def test_cross_page_fragment_offsets(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'source.pdf'
            make_pdf(path)
            kind, units = read_units(path)
        self.assertEqual(kind, 'pdf')
        status, fragments = locate_chunk(units, 'alpha\n\nsecond page')
        self.assertEqual(status, 'exact')
        self.assertEqual([f['unit'] for f in fragments], [1, 2])
        self.assertEqual([f['text'] for f in fragments], ['alpha', 'secondpage'])

    def test_real_word_preserves_paragraph_number_with_empty_paragraph(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'source.docx'
            doc = WordDocument()
            doc.add_paragraph('前文')
            doc.add_paragraph('')
            doc.add_paragraph('引用内容 Unicode ＡＢＣ')
            doc.save(path)
            kind, units = read_units(path)
            status, fragments = locate_chunk(units, '引用内容 Unicode ABC')
        self.assertEqual(kind, 'docx')
        self.assertEqual(status, 'exact')
        self.assertEqual(fragments[0]['unit'], 3)

    def test_repeated_text_requires_offset_and_checks_it(self):
        units = [{'unit': 1, 'text': 'same'}, {'unit': 2, 'text': 'same'}]
        self.assertEqual(locate_chunk(units, 'same')[0], 'ambiguous')
        status, fragments = locate_chunk(units, 'same', 5)
        self.assertEqual(status, 'exact')
        self.assertEqual(fragments[0]['unit'], 2)
        self.assertEqual(fragments[0]['occurrence'], 1)
        self.assertEqual(locate_chunk(units, 'same', 3)[0], 'ambiguous')
        self.assertEqual(locate_chunk(units, 'missing')[0], 'unavailable')

    def test_new_chunks_offsets_match_original_including_overlap(self):
        text = '\n'.join(f'段落{i}，测试引用定位重复内容。' for i in range(200))
        for chunk in chunk_document(text, 'source.docx', 7):
            start = chunk['metadata']['start_index']
            self.assertGreaterEqual(start, 0)
            self.assertEqual(text[start:start + len(chunk['content'])], chunk['content'])

    def test_resolver_source_hash_missing_chunk_and_path_boundary(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'source.docx'
            doc = WordDocument()
            doc.add_paragraph('target paragraph')
            doc.save(path)
            source = Document(id=7, filename='source.docx', file_path=str(path),
                              file_md5=hashlib.md5(path.read_bytes()).hexdigest(), type='knowledge')
            collection = Mock()
            collection.get.return_value = {'ids': ['doc_7_chunk_0'], 'documents': ['target paragraph'],
                                           'metadatas': [{'document_id': '7', 'chunk_index': '0'}]}
            with patch('app.services.vectordb.get_collection', return_value=collection), \
                 patch('app.services.citations.settings.upload_dir', directory):
                response = resolve_citation(source, 0)
                self.assertEqual(response['status'], 'exact')
                self.assertEqual(response['file_type'], 'docx')
                with self.assertRaises(LookupError):
                    resolve_citation(source, -1)
                source.file_md5 = '0' * 32
                with self.assertRaisesRegex(ValueError, '源文件已变化'):
                    resolve_citation(source, 0)
                with patch('app.services.citations.settings.upload_dir', str(Path(directory) / 'restricted')):
                    with self.assertRaisesRegex(ValueError, '访问被拒绝'):
                        resolve_citation(source, 0)
                collection.get.return_value = {'ids': [], 'documents': []}
                with self.assertRaises(LookupError):
                    resolve_citation(source, 0)

    def test_session_document_resolves_without_a_vector_index_entry(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'session.docx'
            doc = WordDocument()
            doc.add_paragraph('session citation target')
            doc.save(path)
            source = Document(
                id=8,
                filename='session.docx',
                file_path=str(path),
                file_md5=hashlib.md5(path.read_bytes()).hexdigest(),
                type='session',
            )
            collection = Mock()
            collection.get.return_value = {'ids': [], 'documents': []}

            with patch('app.services.vectordb.get_collection', return_value=collection), \
                 patch('app.services.citations.settings.upload_dir', directory):
                response = resolve_citation(source, 0)

            self.assertEqual(response['status'], 'exact')
            self.assertEqual(response['text'], 'session citation target')
